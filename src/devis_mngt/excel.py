from __future__ import annotations

import os
import re
import uuid
from copy import copy
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Batch, HEADERS, TRADES, normalized, reasons, unit

TRACE_HEADERS = HEADERS + [
    "Fichier source", "SHA256", "Pages", "Désignation brute", "Unité brute",
    "Extrait source", "Quantité", "Total ligne HT", "Option / variante",
    "Classification", "Motif classification", "Justificatif métadonnées",
    "Statut", "Motifs de vérification", "Ligne Base de prix",
]
REVIEW_HEADERS = TRACE_HEADERS
IGNORED_HEADERS = ["Fichier source", "SHA256", "Statut", "Motif"]
TRACE_NUMERIC = {TRACE_HEADERS.index(header) + 1 for header in (
    "Prix unitaire €HT", "Année du devis", "Quantité", "Total ligne HT",
    "Ligne Base de prix",
)}


def text_cell(cell, value) -> None:
    # openpyxl infers "=" as a formula unless data_type is forced AFTER assignment.
    value = "" if value is None else str(value)
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "\ufffd", value)
    if len(value) > 32767:
        value = value[:32680] + " […] Extrait tronqué pour Excel; texte intégral dans le cache local."
    cell.value = value
    cell.data_type = "s"


def append(sheet, values: list, numeric: set[int] | None = None) -> int:
    row = sheet.max_row + 1
    for index, value in enumerate(values, 1):
        cell = sheet.cell(row, index)
        if numeric and index in numeric and value is not None:
            cell.value = value
        else:
            text_cell(cell, value)
    return row


def sheet_with_headers(workbook, name: str, headers: list[str]):
    sheet = workbook.create_sheet(name)
    for col, header in enumerate(headers, 1):
        text_cell(sheet.cell(1, col), header)
        sheet.cell(1, col).font = Font(bold=True, color="FFFFFF")
        sheet.cell(1, col).fill = PatternFill("solid", fgColor="24476A")
    return sheet


def base_values(extraction, line) -> list:
    designation = line.designation
    if line.option_variant:
        designation += " — " + line.option_variant
    return [line.trade, designation, unit(line.raw_unit),
            float(Decimal(line.unit_price_ht)) if line.unit_price_ht is not None else None,
            extraction.company, extraction.client, extraction.year]


def identity(extraction, line) -> tuple:
    # Never round prices or discard attribution/option semantics when deduplicating.
    return (line.trade, normalized(line.designation), normalized(line.option_variant),
            unit(line.raw_unit), Decimal(line.unit_price_ht),
            normalized(extraction.company), normalized(extraction.client), extraction.year)


def validate_destination(output: Path, template: Path | None = None) -> None:
    if output.suffix.lower() != ".xlsx":
        raise ValueError("La sortie doit avoir l'extension .xlsx")
    if template:
        if output.resolve() == template.resolve():
            raise ValueError("La sortie ne peut pas remplacer le modèle")
        workbook = load_workbook(template)
        try:
            if not any([s.cell(1, i).value for i in range(1, 8)] == HEADERS for s in workbook.worksheets):
                raise ValueError("Le modèle doit contenir les sept en-têtes exacts en A1:G1")
        finally:
            workbook.close()


def export(batch: Batch, output: Path, template: Path | None = None) -> dict:
    validate_destination(output, template)
    protected = [Path(d.source).resolve() for d in batch.documents]
    if output.resolve() in protected or (template and output.resolve() == template.resolve()):
        raise ValueError("La sortie ne peut pas remplacer une source ou le modèle")
    if template:
        workbook = load_workbook(template)
        candidates = [s for s in workbook.worksheets if [s.cell(1, i).value for i in range(1, 8)] == HEADERS]
        if not candidates:
            raise ValueError("Le modèle doit contenir les sept en-têtes exacts en A1:G1")
        base = candidates[0]
        # Only the seven-column input sheet is reused; no unrelated template data.
        for sheet in list(workbook.worksheets):
            if sheet != base:
                workbook.remove(sheet)
        base.title = "Base de prix"
        styles = [copy(base.cell(2, c)._style) for c in range(1, 8)]
        if base.max_row > 1:
            base.delete_rows(2, base.max_row - 1)
        if base.max_column > 7:
            base.delete_cols(8, base.max_column - 7)
        for name in list(base.tables):
            del base.tables[name]
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
        base = sheet_with_headers(workbook, "Base de prix", HEADERS)
        styles = [None] * 7
    trace = sheet_with_headers(workbook, "Traçabilité", TRACE_HEADERS)
    review = sheet_with_headers(workbook, "À vérifier", REVIEW_HEADERS)
    ignored = sheet_with_headers(workbook, "Documents ignorés", IGNORED_HEADERS)
    accepted = {}
    audited = []
    errors = 0
    for document in batch.documents:
        extraction = document.extraction
        if document.status != "extracted" or extraction is None:
            errors += int(document.status == "error")
            append(ignored, [document.source, document.sha256, document.status, document.error])
            audited.append((document, None, [], None))
            continue
        if extraction.classification == "non_devis":
            append(ignored, [document.source, document.sha256, "non_devis", extraction.classification_reason])
        if not extraction.lines:
            audited.append((document, None, [], None))
        for line in extraction.lines:
            flags = reasons(document, line)
            key = None if flags else identity(extraction, line)
            if key is not None:
                accepted.setdefault(key, (extraction, line))
            audited.append((document, line, flags, key))
    mapping = {}
    sorted_items = sorted(accepted.items(), key=lambda item: (
        TRADES.index(item[0][0]), item[0][1], item[0][2], item[0][3],
        item[0][4], item[0][5], item[0][6], item[0][7],
    ))
    for key, (extraction, line) in sorted_items:
        row = append(base, base_values(extraction, line), {4, 7})
        mapping[key] = row
        for col, style in enumerate(styles, 1):
            if style is not None:
                base.cell(row, col)._style = copy(style)
        base.cell(row, 4).number_format = '#,##0.00" €"'
        base.cell(row, 7).number_format = "0"
    for document, line, flags, key in audited:
        extraction = document.extraction
        if line is None:
            values = {
                "Fichier source": document.source,
                "SHA256": document.sha256,
                "Statut": document.status if extraction is None else "sans ligne",
                "Motifs de vérification": document.error if extraction is None else (
                    extraction.classification_reason or "Aucune ligne extraite"
                ),
            }
            append(review, [values.get(header) for header in REVIEW_HEADERS])
            continue
        values = base_values(extraction, line) + [
            document.source, document.sha256, ", ".join(map(str, line.source_pages)),
            line.raw_designation, line.raw_unit, line.source_excerpt,
            float(Decimal(line.quantity)) if line.quantity is not None else None,
            float(Decimal(line.line_total_ht)) if line.line_total_ht is not None else None,
            line.option_variant, extraction.classification, extraction.classification_reason,
            extraction.metadata_excerpt, "À vérifier" if flags else "Accepté",
            "; ".join(flags), mapping.get(key),
        ]
        append(trace, values, TRACE_NUMERIC)
        if flags:
            append(review, values, TRACE_NUMERIC)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index in range(1, sheet.max_column + 1):
            letter = sheet.cell(1, index).column_letter
            if not template or sheet != base or sheet.column_dimensions[letter].width == 13:
                sheet.column_dimensions[letter].width = 55 if index in {2, 8, 13, 18, 20} else 24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        if sheet in (trace, review):
            for row in range(2, sheet.max_row + 1):
                for col in (4, 15):
                    sheet.cell(row, col).number_format = '#,##0.00" €"'
                sheet.cell(row, 7).number_format = "0"
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.{uuid.uuid4().hex}.xlsx")
    try:
        fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        os.close(fd)
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return {"accepted_rows": len(accepted), "extracted_lines": sum(line is not None for _, line, _, _ in audited),
            "review_rows": review.max_row - 1, "errors": errors,
            "status": "ok" if accepted else "aucune donnée acceptée; classeur d'audit uniquement"}
