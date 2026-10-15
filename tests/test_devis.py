"""Offline only. After approval/install: python3 -m unittest discover -s tests -v"""

from __future__ import annotations

import contextlib
import io
import json
import shutil
import unittest
import uuid
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from pypdf import PdfWriter

from devis_mngt.cli import main
from devis_mngt.cloud import AzureBackend, Settings
from devis_mngt.excel import TRACE_HEADERS, export
from devis_mngt.models import Batch, Document, Extraction, HEADERS, Line, OCR, reasons, unit
from devis_mngt.pipeline import process
from devis_mngt.storage import inventory, load, save


def sample(price="25.00", trade="Peinture", designation="Pose de peinture"):
    line = Line(
        trade=trade, designation=designation, raw_designation=designation,
        raw_unit="U", unit_price_ht=price, quantity=None, line_total_ht=None,
        source_pages=[1], source_excerpt=f"{designation} U {price.replace('.', ',')} HT",
        option_variant=None, review_reasons=[],
    )
    metadata = "=Atelier +Client 2025"
    extraction = Extraction(
        classification="devis", classification_reason="Offre de travaux détaillée",
        company="=Atelier", client="+Client", year=2025, metadata_excerpt=metadata,
        filename_conflict=False, review_reasons=[], lines=[line],
    )
    content = metadata + "\n" + line.source_excerpt
    ocr = OCR(content=content, page_texts={"1": content}, page_count=1, raw={})
    return Document(source="devis.pdf", sha256="abc", pages=1, status="extracted",
                    error=None, ocr=ocr, extraction=extraction)


def batch(*documents):
    return Batch(schema_version="1", documents=list(documents))


class LocalFiles(unittest.TestCase):
    def setUp(self):
        # Never use system temporary directories, including in tests.
        self.directory = Path.cwd() / ".test-work" / uuid.uuid4().hex
        self.directory.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.directory)
        try:
            self.directory.parent.rmdir()
        except OSError:
            pass

    def pdf(self, name="devis.pdf", pages=1, encrypted=False):
        path = self.directory / name
        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=72, height=72)
        if encrypted:
            writer.encrypt("test-only-password")
        with path.open("wb") as stream:
            writer.write(stream)
        writer.close()
        return path


class ModelTests(unittest.TestCase):
    def test_supported_units(self):
        for raw, expected in [("m2", "m²"), ("m²", "m²"), ("m3", "m³"),
                              ("m³", "m³"), ("ml", "ml"), ("u", "U"), ("forfait", "Forfait")]:
            self.assertEqual(unit(raw), expected)
        self.assertIsNone(unit("inconnu"))

    def test_missing_quantity_does_not_invent_or_reject(self):
        doc = sample()
        self.assertEqual(reasons(doc, doc.extraction.lines[0]), [])
        self.assertIsNone(doc.extraction.lines[0].quantity)

    def test_unknown_price_is_review(self):
        doc = sample()
        doc.extraction.lines[0].unit_price_ht = None
        self.assertTrue(any("Prix unitaire HT absent" in r for r in reasons(doc, doc.extraction.lines[0])))

    def test_math_mismatch_is_review(self):
        doc = sample()
        line = doc.extraction.lines[0]
        line.quantity, line.line_total_ht = "2", "49.90"
        self.assertTrue(any("Quantité" in r for r in reasons(doc, line)))
        line.line_total_ht = "50.01"
        self.assertEqual(reasons(doc, line), [])

    def test_hallucinated_or_wrong_page_evidence_is_review(self):
        for change in ("source_excerpt", "source_pages", "unit_price_ht", "raw_designation"):
            with self.subTest(change=change):
                doc = sample()
                line = doc.extraction.lines[0]
                setattr(line, change, {
                    "source_excerpt": "Texte inventé", "source_pages": [99],
                    "unit_price_ht": "125.00", "raw_designation": "Désignation inventée",
                }[change])
                self.assertTrue(reasons(doc, line))

    def test_ambiguous_trade_and_metadata_is_review(self):
        doc = sample()
        doc.extraction.company = "Autre entreprise"
        doc.extraction.lines[0].trade = "Autres"
        flags = reasons(doc, doc.extraction.lines[0])
        self.assertTrue(any("Entreprise" in r for r in flags))
        self.assertTrue(any("Corps d'état" in r for r in flags))

    def test_uncertain_quote_keeps_lines_for_review(self):
        doc = sample()
        doc.extraction.classification = "incertain"
        self.assertEqual(len(doc.extraction.lines), 1)
        self.assertTrue(reasons(doc, doc.extraction.lines[0]))

    def test_decimal_validation_is_closed(self):
        for bad in ("NaN", "Infinity", "12,50", "1e9", "1" * 40):
            with self.subTest(bad=bad), self.assertRaises(ValidationError):
                Line.model_validate({**sample().extraction.lines[0].model_dump(), "unit_price_ht": bad})
        with self.assertRaises(ValidationError):
            Extraction.model_validate({**sample().extraction.model_dump(), "unexpected": True})

    def test_schema_all_properties_required_closed(self):
        schema = Extraction.model_json_schema()
        self.assertFalse(schema["additionalProperties"])
        self.assertEqual(set(schema["required"]), set(schema["properties"]))
        nested = schema["$defs"]["Line"]
        self.assertFalse(nested["additionalProperties"])
        self.assertEqual(set(nested["required"]), set(nested["properties"]))


class ExcelTests(LocalFiles):
    def test_exact_headers_numeric_cells_formula_safety_and_duplicates(self):
        first, duplicate, different = sample(), sample(), sample("25.01")
        duplicate.source = "other/devis.pdf"
        different.source = "third.pdf"
        file = self.directory / "prix.xlsx"
        summary = export(batch(first, duplicate, different), file)
        self.assertEqual(summary["accepted_rows"], 2)
        book = load_workbook(file)
        try:
            base, trace = book["Base de prix"], book["Traçabilité"]
            self.assertEqual([c.value for c in base[1]], HEADERS)
            self.assertEqual(base.max_row, 3)
            self.assertEqual(trace.max_row, 4)
            self.assertEqual(base.cell(2, 4).data_type, "n")
            self.assertEqual(base.cell(2, 7).data_type, "n")
            self.assertEqual(base.cell(2, 5).value, "=Atelier")
            self.assertEqual(base.cell(2, 5).data_type, "s")
            self.assertEqual(base.cell(2, 6).value, "+Client")
            self.assertEqual(base.cell(2, 6).data_type, "s")
            base_column = TRACE_HEADERS.index("Ligne Base de prix") + 1
            self.assertEqual(trace.cell(2, base_column).value, trace.cell(3, base_column).value)
            self.assertNotEqual(trace.cell(2, base_column).value, trace.cell(4, base_column).value)
            self.assertEqual(trace.cell(2, base_column).data_type, "n")
            self.assertEqual(base.freeze_panes, "A2")
            self.assertEqual(base.auto_filter.ref, "A1:G3")
            self.assertFalse(any(c.data_type == "f" for s in book for row in s for c in row))
        finally:
            book.close()

    def test_no_dedup_across_company_client_year_or_options(self):
        docs = []
        for i, change in enumerate(("none", "company", "client", "year", "option")):
            doc = sample()
            if change == "company":
                doc.extraction.company = "Autre entreprise"
            elif change == "client":
                doc.extraction.client = "Autre client"
            elif change == "year":
                doc.extraction.year = 2024
            elif change == "option":
                doc.extraction.lines[0].option_variant = "Option finition"
            metadata = f"{doc.extraction.company} {doc.extraction.client} {doc.extraction.year}"
            doc.extraction.metadata_excerpt = metadata
            doc.ocr.content = metadata + "\n" + doc.extraction.lines[0].source_excerpt
            doc.ocr.page_texts = {"1": doc.ocr.content}
            doc.source = f"source-{i}.pdf"
            docs.append(doc)
        result = export(batch(*docs), self.directory / "prix.xlsx")
        self.assertEqual(result["accepted_rows"], 5)

    def test_trade_order_not_alphabetical(self):
        docs = [sample(trade=t) for t in ("Peinture", "Électricité", "Démolition", "Couverture")]
        file = self.directory / "prix.xlsx"
        export(batch(*docs), file)
        book = load_workbook(file)
        try:
            self.assertEqual([r[0].value for r in list(book["Base de prix"])[1:]],
                             ["Démolition", "Couverture", "Électricité", "Peinture"])
        finally:
            book.close()

    def test_template_styles_and_no_example_data_copied(self):
        template = self.directory / "vierge.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Feuil1"
        sheet.append(HEADERS)
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor="FFAA00")
        sheet.cell(2, 2).font = Font(name="Arial", italic=True)
        sheet.cell(2, 2).value = "NE PAS COPIER"
        sheet.column_dimensions["B"].width = 66
        book.save(template)
        book.close()
        output = self.directory / "prix.xlsx"
        export(batch(sample()), output, template)
        result = load_workbook(output)
        try:
            sheet = result["Base de prix"]
            self.assertEqual(sheet.cell(1, 1).fill.fgColor.rgb, "00FFAA00")
            self.assertTrue(sheet.cell(2, 2).font.italic)
            self.assertEqual(sheet.column_dimensions["B"].width, 66)
            self.assertEqual(sheet.cell(2, 2).value, "Pose de peinture")
        finally:
            result.close()

    def test_missing_price_error_unsupported_and_nonquote_audited(self):
        unknown = sample()
        unknown.extraction.lines[0].unit_price_ht = None
        error = Document(source="corrupt.pdf", sha256="def", pages=None, status="error",
                         error="PDF illisible", ocr=None, extraction=None)
        unsupported = error.model_copy(update={"source": "notes.txt", "status": "unsupported",
                                                "error": "Format non pris en charge"})
        nonquote = sample()
        nonquote.extraction.classification = "non_devis"
        file = self.directory / "prix.xlsx"
        summary = export(batch(unknown, error, unsupported, nonquote), file)
        self.assertEqual(summary["accepted_rows"], 0)
        self.assertEqual(summary["errors"], 1)
        self.assertIn("audit", summary["status"])
        book = load_workbook(file)
        try:
            self.assertEqual(book["Base de prix"].max_row, 1)
            self.assertEqual(book["À vérifier"].max_row, 5)
            self.assertEqual(book["Traçabilité"].max_row, 3)
            self.assertEqual(book["Documents ignorés"].max_row, 4)
            reasons_column = TRACE_HEADERS.index("Motifs de vérification") + 1
            status_column = TRACE_HEADERS.index("Statut") + 1
            self.assertEqual(book["À vérifier"].cell(3, reasons_column).value, "PDF illisible")
            self.assertEqual(book["À vérifier"].cell(3, status_column).value, "error")
        finally:
            book.close()

    def test_all_text_prefixes_are_strings(self):
        for prefix in ("=", "+", "-", "@"):
            doc = sample(designation=prefix + "ouvrage")
            doc.source = prefix + "devis.pdf"
            output = self.directory / f"prix-{ord(prefix)}.xlsx"
            export(batch(doc), output)
            book = load_workbook(output)
            try:
                self.assertEqual(book["Base de prix"].cell(2, 2).data_type, "s")
                self.assertEqual(book["Base de prix"].cell(2, 2).value, prefix + "ouvrage")
                self.assertEqual(book["Traçabilité"].cell(2, 8).data_type, "s")
            finally:
                book.close()

    def test_original_and_template_cannot_be_overwritten(self):
        file = self.directory / "source.xlsx"
        doc = sample()
        doc.source = str(file)
        with self.assertRaises(ValueError):
            export(batch(doc), file)


class StorageAndPipelineTests(LocalFiles):
    settings = Settings("https://di.example", "https://aoai.example", "test-deployment")

    def test_inventory_pages_corrupt_encrypted_unsupported_and_sha(self):
        good = self.pdf(pages=2)
        self.pdf("encrypted.pdf", encrypted=True)
        (self.directory / "corrupt.pdf").write_bytes(b"not a PDF")
        (self.directory / "notes.txt").write_text("not supported", encoding="utf-8")
        rows = {Path(r["source"]).name: r for r in inventory(self.directory)}
        self.assertEqual(rows[good.name]["pages"], 2)
        self.assertEqual(len(rows[good.name]["sha256"]), 64)
        self.assertEqual(rows["encrypted.pdf"]["status"], "error")
        self.assertEqual(rows["corrupt.pdf"]["status"], "error")
        self.assertEqual(rows["notes.txt"]["status"], "unsupported")
        self.assertEqual(len(inventory(self.directory, 1)), 1)

    def test_cache_atomic_checksum_and_schema(self):
        file = self.directory / "cache.json"
        save(file, sample().model_dump())
        self.assertEqual(Document.model_validate(load(file)).source, "devis.pdf")
        envelope = json.loads(file.read_text())
        envelope["payload"]["source"] = "edited.pdf"
        file.write_text(json.dumps(envelope))
        with self.assertRaises(ValueError):
            load(file)
        self.assertFalse(list(self.directory.glob("*.part")))

    def test_duplicates_cached_and_original_names_preserved(self):
        folder = self.directory / "input"
        folder.mkdir()
        original = self.pdf()
        shutil.copyfile(original, folder / "a.pdf")
        shutil.copyfile(original, folder / "b.pdf")
        backend = MagicMock()
        backend.ocr.return_value = sample().ocr
        backend.extract.return_value = sample().extraction
        cache = self.directory / "cache"
        first = process(folder, cache, self.settings, backend=backend)
        second = process(folder, cache, self.settings, backend=backend)
        self.assertEqual(backend.ocr.call_count, 1)
        self.assertEqual(backend.extract.call_count, 1)
        self.assertEqual([Path(d.source).name for d in first.documents], ["a.pdf", "b.pdf"])
        self.assertTrue(all(d.status == "extracted" for d in second.documents))
        summary = export(first, self.directory / "prices.xlsx")
        self.assertEqual(summary["accepted_rows"], 1)
        self.assertEqual(summary["extracted_lines"], 2)

    def test_configuration_change_only_repeats_affected_stage(self):
        source = self.pdf()
        backend = MagicMock()
        backend.ocr.return_value = sample().ocr
        backend.extract.return_value = sample().extraction
        cache = self.directory / "cache"
        process(source, cache, self.settings, backend=backend)
        newer = Settings(self.settings.di_endpoint, self.settings.openai_endpoint,
                         self.settings.deployment, model_revision="2")
        process(source, cache, newer, backend=backend)
        self.assertEqual(backend.ocr.call_count, 1)
        self.assertEqual(backend.extract.call_count, 2)

    def test_failed_extraction_resumes_ocr_and_reports_error(self):
        source = self.pdf()
        backend = MagicMock()
        backend.ocr.return_value = sample().ocr
        backend.extract.side_effect = RuntimeError("service unavailable")
        cache = self.directory / "cache"
        with contextlib.redirect_stderr(io.StringIO()) as stderr:
            failed = process(source, cache, self.settings, backend=backend)
        self.assertEqual(failed.documents[0].status, "error")
        self.assertIn(str(source), stderr.getvalue())
        self.assertIsNotNone(failed.documents[0].ocr)
        backend.extract.side_effect = None
        backend.extract.return_value = sample().extraction
        success = process(source, cache, self.settings, backend=backend)
        self.assertEqual(success.documents[0].status, "extracted")
        self.assertEqual(backend.ocr.call_count, 1)

    def test_incomplete_ocr_rejected_before_llm(self):
        source = self.pdf(pages=4)
        backend = MagicMock()
        backend.ocr.return_value = sample().ocr
        with contextlib.redirect_stderr(io.StringIO()):
            result = process(source, self.directory / "cache", self.settings, backend=backend)
        self.assertEqual(result.documents[0].status, "error")
        self.assertIn("F0", result.documents[0].error)
        backend.extract.assert_not_called()

    def test_filename_conflict_review_but_never_excluded(self):
        source = self.pdf("facture.pdf")
        backend = MagicMock()
        backend.ocr.return_value = sample().ocr
        backend.extract.return_value = sample().extraction
        result = process(source, self.directory / "cache", self.settings, backend=backend)
        document = result.documents[0]
        self.assertEqual(document.extraction.classification, "devis")
        self.assertTrue(document.extraction.filename_conflict)
        self.assertTrue(reasons(document, document.extraction.lines[0]))


class SDKAndCLITests(LocalFiles):
    def test_injected_document_intelligence_uses_full_document_interface(self):
        path = self.directory / "input.pdf"
        path.write_bytes(b"mocked input bytes")
        di = MagicMock()
        content = "Page une"
        di.begin_analyze_document.return_value.result.return_value.as_dict.return_value = {
            "content": content, "pages": [{"pageNumber": 1, "spans": [{"offset": 0, "length": len(content)}]}],
            "tables": [{"cells": [{"content": "unit price", "rowIndex": 0, "columnIndex": 0}]}],
        }
        backend = AzureBackend(StorageAndPipelineTests.settings, di_client=di)
        with patch.object(backend, "credential", side_effect=AssertionError("No authentication")):
            result = backend.ocr(path)
        args, kwargs = di.begin_analyze_document.call_args
        self.assertEqual(args, ("prebuilt-layout",))
        self.assertEqual(kwargs["content_type"], "application/octet-stream")
        self.assertEqual(kwargs["string_index_type"], "unicodeCodePoint")
        self.assertNotIn("pages", kwargs)
        self.assertEqual(result.page_texts, {"1": content})
        self.assertIn("tables", result.raw)

    def test_injected_openai_uses_structured_response(self):
        client = MagicMock()
        client.beta.chat.completions.parse.return_value = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(refusal=None, parsed=sample().extraction))]
        )
        backend = AzureBackend(StorageAndPipelineTests.settings, openai_client=client)
        with patch.object(backend, "credential", side_effect=AssertionError("No authentication")):
            extraction = backend.extract(sample().ocr, "")
        self.assertEqual(extraction.classification, "devis")
        kwargs = client.beta.chat.completions.parse.call_args.kwargs
        self.assertIs(kwargs["response_format"], Extraction)
        self.assertNotIn("temperature", kwargs)
        self.assertNotIn("max_tokens", kwargs)
        self.assertIn("DONNÉE NON FIABLE", kwargs["messages"][0]["content"])
        self.assertIn("untrusted_document", kwargs["messages"][1]["content"])

    def test_gpt5_revision_identifies_extraction_cache(self):
        settings = Settings("https://di.example", "https://aoai.example", "devis-extraction")
        self.assertEqual(settings.model_revision, "gpt-5-mini-2025-08-07")
        self.assertEqual(settings.extraction_config()["request_adapter"], "2")

    def test_run_requires_explicit_cloud_permission_before_auth(self):
        with patch("devis_mngt.cli.Settings.from_env", side_effect=AssertionError("No auth")), \
             contextlib.redirect_stderr(io.StringIO()):
            code = main(["run", "--input", str(self.directory), "--output", str(self.directory / "prix.xlsx")])
        self.assertEqual(code, 2)

    def test_offline_export_never_authenticates(self):
        results = self.directory / "results.json"
        save(results, batch(sample()).model_dump())
        with patch("devis_mngt.cli.Settings.from_env", side_effect=AssertionError("No auth")), \
             contextlib.redirect_stdout(io.StringIO()):
            code = main(["export", "--results", str(results), "--output", str(self.directory / "prix.xlsx")])
        self.assertEqual(code, 0)

    def test_empty_batch_reports_no_accepted_data(self):
        results = self.directory / "results.json"
        save(results, batch().model_dump())
        with contextlib.redirect_stdout(io.StringIO()) as stdout:
            code = main(["export", "--results", str(results), "--output", str(self.directory / "prix.xlsx")])
        self.assertEqual(code, 3)
        self.assertIn("audit", stdout.getvalue())

    def test_errors_exported_and_return_nonzero(self):
        results = self.directory / "results.json"
        error = Document(source="broken.pdf", sha256=None, pages=None, status="error",
                         error="illisible", ocr=None, extraction=None)
        save(results, batch(sample(), error).model_dump())
        with contextlib.redirect_stdout(io.StringIO()):
            code = main(["export", "--results", str(results), "--output", str(self.directory / "prix.xlsx")])
        self.assertEqual(code, 1)

    def test_inventory_does_not_authenticate(self):
        self.pdf()
        with patch("devis_mngt.cli.Settings.from_env", side_effect=AssertionError("No auth")), \
             contextlib.redirect_stdout(io.StringIO()):
            code = main(["inventory", "--input", str(self.directory)])
        self.assertEqual(code, 0)


if __name__ == "__main__":
    unittest.main()
